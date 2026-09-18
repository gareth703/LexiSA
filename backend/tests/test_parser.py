"""Checkpoint 2.1: exercise PDF and DOCX extraction."""
from __future__ import annotations

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parents[1]))

import fitz
from docx import Document

from services.parser import extract_text_from_file

FIXTURE_DIR = Path(__file__).parent / "fixtures"


def build_pdf() -> bytes:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "MASTER SERVICES AGREEMENT\nPage 1\n\n1.1 Warranty\nThe Provider warrants all deliverables for 30 days from delivery.\n\n2.3 Liability\nThe Provider will maintain reasonable safeguards.")
    return document.tobytes()


def build_docx() -> bytes:
    document = Document()
    document.add_heading("MASTER SERVICES AGREEMENT", level=1)
    document.add_paragraph("1.1 Warranty")
    document.add_paragraph("The Provider warrants all deliverables for 30 days from delivery.")
    document.add_paragraph("2.3 Liability")
    document.add_paragraph("The Provider will maintain reasonable safeguards.")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


if __name__ == "__main__":
    FIXTURE_DIR.mkdir(exist_ok=True)
    pdf_path = FIXTURE_DIR / "sample.pdf"
    docx_path = FIXTURE_DIR / "sample.docx"
    pdf_path.write_bytes(build_pdf())
    docx_path.write_bytes(build_docx())
    pdf_text = extract_text_from_file(pdf_path.read_bytes(), pdf_path.name)
    docx_text = extract_text_from_file(docx_path.read_bytes(), docx_path.name)
    print("PDF extraction:\n", pdf_text)
    print("\nDOCX extraction:\n", docx_text)
    assert "1.1 Warranty" in pdf_text and "2.3 Liability" in pdf_text
    assert "1.1 Warranty" in docx_text and "2.3 Liability" in docx_text
    assert "Page 1" not in pdf_text
    print("Checkpoint 2.1 passed")
